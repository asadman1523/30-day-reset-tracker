import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

import run_tracker_updates as runner  # noqa: E402


HEADERS = ["日期", "餐別", "餐點", "來源", "蛋白質", "碳水", "熱量", "備註"]
DAILY_HEADERS = [
    "日期",
    "零下注",
    "查看盤口",
    "衝動 0–10",
    "觸發原因",
    "替代行動",
    "完成重訓",
    "散步分鐘",
    "睡眠小時",
    "連續零下注",
]


class TrackerFoodUpdateTests(unittest.TestCase):
    def make_food_sheet(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "飲食紀錄"
        ws.append(HEADERS)
        return ws

    def add_old_combined_breakfast(self, ws):
        ws.append(
            [
                date(2026, 9, 7),
                "早餐",
                "地瓜＋雞胸肉 1 份",
                "使用者回報",
                "",
                "",
                "",
                "原始合併紀錄",
            ]
        )

    def sweet_potato_update(self):
        return {
            "type": "food",
            "operation_id": "test-sweet-potato",
            "date": "2026-09-07",
            "meal": "早餐",
            "match_food": "地瓜＋雞胸肉 1 份",
            "food": "7-ELEVEN 30元地瓜",
            "source": "測試估算",
            "protein": 2.3,
            "carbs": 37.7,
            "calories": 162,
            "note": "地瓜獨立紀錄",
            "update_existing": True,
        }

    def chicken_update(self):
        return {
            "type": "food",
            "operation_id": "test-chicken",
            "date": "2026-09-07",
            "meal": "早餐",
            "food": "7-ELEVEN 21Plus 剝皮辣椒雞胸肉 1 包（145g）",
            "source": "既有同款紀錄",
            "protein": 30,
            "calories": 163,
            "note_append": "沿用先前同款營養資料。",
        }

    def test_match_food_can_rename_existing_row_without_duplicate(self):
        ws = self.make_food_sheet()
        self.add_old_combined_breakfast(ws)

        result = runner.apply_food(ws, self.sweet_potato_update())

        self.assertEqual(result["status"], "updated")
        self.assertEqual(result["row"], 2)
        self.assertEqual(ws.max_row, 2)
        self.assertEqual(ws.cell(2, 3).value, "7-ELEVEN 30元地瓜")
        self.assertIn("food", result["fields"])
        verification = runner.verify_food(ws, result)
        self.assertTrue(verification["verified"])
        self.assertEqual(verification["row"], 2)

    def test_note_append_is_written_when_food_is_new(self):
        ws = self.make_food_sheet()

        result = runner.apply_food(ws, self.chicken_update())

        self.assertEqual(result["status"], "written")
        self.assertEqual(result["row"], 2)
        self.assertIn("沿用先前同款營養資料。", str(ws.cell(2, 8).value))
        verification = runner.verify_food(ws, result)
        self.assertTrue(verification["verified"])

    def test_combined_breakfast_can_be_split_into_two_independent_rows(self):
        ws = self.make_food_sheet()
        self.add_old_combined_breakfast(ws)

        sweet_result = runner.apply_food(ws, self.sweet_potato_update())
        chicken_result = runner.apply_food(ws, self.chicken_update())

        self.assertTrue(runner.verify_food(ws, sweet_result)["verified"])
        self.assertTrue(runner.verify_food(ws, chicken_result)["verified"])
        self.assertEqual(ws.max_row, 3)
        foods = [ws.cell(row, 3).value for row in range(2, 4)]
        self.assertEqual(
            foods,
            [
                "7-ELEVEN 30元地瓜",
                "7-ELEVEN 21Plus 剝皮辣椒雞胸肉 1 包（145g）",
            ],
        )
        proteins = [ws.cell(row, 5).value for row in range(2, 4)]
        self.assertAlmostEqual(sum(proteins), 32.3)
        self.assertEqual(ws.cell(2, 6).value, 37.7)
        self.assertIn(ws.cell(3, 6).value, (None, ""))

    def test_processed_operation_ids_are_deduplicated(self):
        original_path = runner.PROCESSED_IDS
        try:
            with tempfile.TemporaryDirectory() as tmpdir:
                runner.PROCESSED_IDS = Path(tmpdir) / "processed_operation_ids.json"
                runner._write_processed_ids(["op-a", "op-a", "op-b"])
                self.assertEqual(runner._load_processed_ids(), ["op-a", "op-b"])
        finally:
            runner.PROCESSED_IDS = original_path


class TrackerDailyCheckinTests(unittest.TestCase):
    def make_daily_sheet(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "每日檢核"
        ws.append(DAILY_HEADERS)
        ws.append([date(2026, 9, 2), None, None, None, None, None, None, 10, None, "=1"])
        return ws

    def test_missing_date_is_created_and_walk_minutes_written(self):
        ws = self.make_daily_sheet()
        item = {
            "type": "daily_checkin",
            "date": "2026-09-11",
            "walk_minutes": 20,
        }

        result = runner.apply_daily_checkin(ws, item)

        self.assertEqual(result["status"], "written")
        self.assertTrue(result["created"])
        self.assertEqual(ws.max_row, 3)
        self.assertEqual(ws.cell(3, 1).value, date(2026, 9, 11))
        self.assertEqual(ws.cell(3, 8).value, 20)
        self.assertTrue(runner.verify_daily_checkin(ws, result)["verified"])

    def test_existing_date_is_updated_without_duplicate_row(self):
        ws = self.make_daily_sheet()
        item = {
            "type": "daily_checkin",
            "date": "2026-09-02",
            "walk_minutes": 25,
        }

        result = runner.apply_daily_checkin(ws, item)

        self.assertEqual(result["status"], "updated")
        self.assertFalse(result["created"])
        self.assertEqual(ws.max_row, 2)
        self.assertEqual(ws.cell(2, 8).value, 25)
        self.assertTrue(runner.verify_daily_checkin(ws, result)["verified"])

    def test_new_date_only_writes_fields_explicitly_provided(self):
        ws = self.make_daily_sheet()
        item = {
            "type": "daily_checkin",
            "date": "2026-09-11",
            "walk_minutes": 20,
        }

        result = runner.apply_daily_checkin(ws, item)
        row = result["row"]

        for col in (2, 3, 4, 5, 6, 7, 9, 10):
            self.assertIn(ws.cell(row, col).value, (None, ""))
        self.assertEqual(ws.cell(row, 8).value, 20)


if __name__ == "__main__":
    unittest.main()
